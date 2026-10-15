"""样本工作流：读 docs/sample 的源表与已填模板，调用对齐模块，把 TOML 打到标准输出。"""

from __future__ import annotations

import sys
import time
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from geodesic_align.digit_processor import TableView, _excel_col, id_lineage
from geodesic_align.emit_mapping import (
    MappingDecision,
    build_sources,
    recover_mapping,
    regex_toml_literal,
    resolve_text_derived_morphism,
)
from geodesic_align.text_processor import run_text_align


_EMPTY = frozenset({"", "null", "none", "n/a", "na"})
_SAMPLE_DIR = Path(__file__).resolve().parents[1] / "docs" / "sample"
_SOURCE_BOOK = "执法堂业绩.xlsx"
_DEST_BOOK = "sample_template.xlsx"
_SOURCE_SHEETS = ("sheet1", "sheet2")
_DEST_SHEET = "Input_sheet"
_PRINT_SHEET = "Print_sheet"



def is_empty(text: str) -> bool:
    """
    函数名: is_empty
    作用: 判断格子字面是否为空
    输入:
        text (str): 单元格字面
    输出:
        bool
    """
    return str(text).strip().lower() in _EMPTY



def cell_text(value: object) -> str:
    """
    函数名: cell_text
    作用: 单元格转字面；datetime 保留时刻，date 只留日期
    输入:
        value (object): openpyxl 单元格值
    输出:
        str
    """
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.hour or value.minute or value.second:
            return value.strftime("%Y-%m-%d %H:%M:%S")
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()



def load_header_rows(path: Path, sheet: str, *, header_row: int = 1) -> tuple[list[str], list[dict[str, str]]]:
    """
    函数名: load_header_rows
    作用: 读一张工作表的表头与数据行；公式格保留以 = 开头的原文
    输入:
        path (Path): 工作簿路径
        sheet (str): 工作表名
        header_row (int): 表头行，1-based
    输出:
        tuple[list[str], list[dict[str, str]]]
    """
    # 中文注释: data_only=False 才能留下以 = 开头的公式原型
    book = load_workbook(path, data_only=False)
    try:
        if sheet not in book.sheetnames:
            raise KeyError(sheet)
        ws = book[sheet]
        max_col = int(ws.max_column or 1)
        max_row = int(ws.max_row or 1)
        headers: list[str] = []
        for col in range(1, max_col + 1):
            name = cell_text(ws.cell(int(header_row), col).value)
            if name:
                headers.append(name)
        rows: list[dict[str, str]] = []
        for row_idx in range(int(header_row) + 1, max_row + 1):
            item: dict[str, str] = {}
            for col, name in enumerate(headers, start=1):
                item[name] = cell_text(ws.cell(row_idx, col).value)
            if any(not is_empty(text) for text in item.values()):
                rows.append(item)
        return headers, rows
    finally:
        book.close()



def sheet_names(path: Path) -> list[str]:
    """
    函数名: sheet_names
    作用: 列出工作簿里的工作表名
    输入:
        path (Path): 工作簿路径
    输出:
        list[str]
    """
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        return list(book.sheetnames)
    finally:
        book.close()



def excel_locator(path: Path, sheet: str) -> str:
    """
    函数名: excel_locator
    作用: 写成 Excel 外链字面：目录\\[文件名]工作表
    输入:
        path (Path): 工作簿
        sheet (str): 工作表名
    输出:
        str
    """
    return str(path.parent) + "\\[" + path.name + "]" + sheet



def _toml_basic(text: str) -> str:
    """
    函数名: _toml_basic
    作用: 把普通字符串收成 TOML 双引号字面
    输入:
        text (str): 原文
    输出:
        str
    """
    escaped = str(text).replace("\\", "\\\\").replace('"', '\\"')
    return f'"{escaped}"'



def load_sample_views() -> tuple[Path, Path, list[TableView], TableView]:
    """
    函数名: load_sample_views
    作用: 读执法堂业绩两张源表和 sample_template 的 Input_sheet
    输入: 无
    输出:
        tuple: (源工作簿, 模板工作簿, 源表, 模板表)
    """
    source_path = _SAMPLE_DIR / _SOURCE_BOOK
    dest_path = _SAMPLE_DIR / _DEST_BOOK
    if not source_path.is_file() or not dest_path.is_file():
        raise FileNotFoundError(f"样本工作簿不在 {_SAMPLE_DIR}")
    sources: list[TableView] = []
    for idx, sheet in enumerate(_SOURCE_SHEETS, start=1):
        headers, rows = load_header_rows(source_path, sheet)
        sources.append(TableView(f"source{idx}", "excel", headers, rows, sheet=sheet, header_row=1))
    headers, rows = load_header_rows(dest_path, _DEST_SHEET)
    dest = TableView("dest", "excel", headers, rows, sheet=_DEST_SHEET, header_row=1)
    return source_path, dest_path, sources, dest



def _source_key(sources: list[TableView], sheet: str) -> str:
    """
    函数名: _source_key
    作用: 把测地线里的工作表名换成 [sources] 键
    输入:
        sources (list[TableView]): 源表
        sheet (str): 工作表名或已经是 source 键
    输出:
        str
    """
    for table in sources:
        if table.sheet == sheet or table.source_key == sheet:
            return table.source_key
    return sheet



def _apply_text_end(
    dest: TableView,
    sources: list[TableView],
    lineage: list[list[tuple[str, int]]],
    decision: MappingDecision,
    end: object,
) -> None:
    """
    函数名: _apply_text_end
    作用: 只把仍是手工输入的列改成测地线给出的搬运或拼接
    输入:
        dest (TableView): 模板
        sources (list[TableView]): 源表
        lineage (list): 行指针
        decision (MappingDecision): 当前仍为 input 的列
        end (object): GeodesicEnd
    输出: 无
    """
    sheets = [bit for bit in str(getattr(end, "source_sheet", "")).split("|") if bit]
    names = [bit for bit in str(getattr(end, "field", "")).split("|") if bit]
    klass = str(getattr(end, "klass", ""))
    if klass == "carry" and sheets and names:
        decision.cell_role = "data"
        decision.source = _source_key(sources, sheets[0])
        decision.field = names[0]
        decision.klass = klass
        decision.confidence = 1.0
        return
    if klass == "derived" and len(names) >= 2:
        if len(sheets) == len(names):
            pairs = [(_source_key(sources, sheet), name) for sheet, name in zip(sheets, names)]
        elif len(sheets) == 1:
            key = _source_key(sources, sheets[0])
            pairs = [(key, name) for name in names]
        else:
            pairs = list(zip(sheets, names))
        role, source, field_name, formula, warnings = resolve_text_derived_morphism(
            dest, decision.input_label, sources, lineage, pairs,
        )
        decision.cell_role = role
        decision.source = source
        decision.field = field_name
        decision.formula = formula
        decision.klass = klass
        decision.confidence = 0.5 if role == "formula" else 0.0
        decision.warnings.extend(warnings)



def align_sample(
    sources: list[TableView],
    dest: TableView,
) -> tuple[list[MappingDecision], dict[str, float]]:
    """
    函数名: align_sample
    作用: 数字、正则和行指针先定案，测地线只补仍是手工输入的列
    输入:
        sources (list[TableView]): 源表
        dest (TableView): 已填模板
    输出:
        tuple[list[MappingDecision], dict[str, float]]: 决策，以及 map / text 两段毫秒数
    """
    mark = time.perf_counter()
    decisions = recover_mapping(sources, dest)
    map_ms = (time.perf_counter() - mark) * 1000
    # 中文注释: 测地线含 BERT 前向，单独计时
    mark = time.perf_counter()
    text_ends = run_text_align(
        [(table.sheet, table.headers, table.rows) for table in sources],
        dest.sheet,
        dest.headers,
        dest.rows,
    )
    by_label = {item.input_label: item for item in decisions}
    lineage = id_lineage(dest, sources)
    for end in text_ends:
        current = by_label.get(str(getattr(end, "input_label", "")))
        if current is None or current.cell_role != "input":
            continue
        _apply_text_end(dest, sources, lineage, current, end)
    text_ms = (time.perf_counter() - mark) * 1000
    return decisions, {"map": map_ms, "text": text_ms}



def render_toml(
    source_path: Path,
    dest_path: Path,
    sources: list[TableView],
    dest: TableView,
    decisions: list[MappingDecision],
) -> str:
    """
    函数名: render_toml
    作用: 把源定位和映射决策收成模板 TOML
    输入:
        source_path (Path): 源工作簿
        dest_path (Path): 模板工作簿
        sources (list[TableView]): 源表，顺序与 source1、source2 一致
        dest (TableView): 模板表
        decisions (list[MappingDecision]): 每条模板列一条
    输出:
        str: TOML 文本
    """
    locators = {
        table.source_key: excel_locator(source_path, table.sheet)
        for table in sources
    }
    build_sources(locators)
    names = sheet_names(dest_path)
    print_sheet = _PRINT_SHEET if _PRINT_SHEET in names else ""
    last_col = _excel_col(max(len(dest.headers), 1))
    lines = [
        "determiner = \"\\t\"",
        f"work_sheet = {_toml_basic(dest.sheet)}",
        f"print_sheet = {_toml_basic(print_sheet)}",
        "use_independent_db = true",
    ]
    id_labels = [item.input_label for item in decisions if item.is_id]
    if len(id_labels) == 1:
        lines.append(f"db_id = {_toml_basic(id_labels[0])}")
    lines.append("")
    lines.append("[sources]")
    for key, locator in locators.items():
        lines.append(f"{key} = '{locator}'")
    lines.append("")
    lines.append("[[input_section]]")
    lines.append(f'input_area = "A2:{last_col}2"')
    lines.append('move_to = "down"')
    lines.append("offset = 1")
    lines.append("")
    paste_index = 0
    for item in decisions:
        for warning in item.warnings:
            lines.append("# " + warning.replace("\n", " "))
        lines.append("[[fields]]")
        lines.append(f"Input_label = {_toml_basic(item.input_label)}")
        lines.append('value_from_label = "down"')
        lines.append("value_offset = 1")
        lines.append(f"cell_role = {_toml_basic(item.cell_role)}")
        lines.append(f"source = {_toml_basic(item.source)}")
        lines.append(f"field = {_toml_basic(item.field)}")
        if item.cell_role == "regex" and item.regex:
            lines.append(regex_toml_literal(item.regex))
        else:
            lines.append("regex = ''")
        lines.append(f"formula = {_toml_basic(item.formula)}")
        if item.cell_role in {"data", "regex"}:
            index = paste_index
            paste_index += 1
        else:
            index = -1
        lines.append(f"index = {index}")
        lines.append("id = true" if item.is_id else "id = false")
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"



def build_sample_toml() -> str:
    """
    函数名: build_sample_toml
    作用: 读样本工作簿，对齐后得到 TOML 文本
    输入: 无
    输出:
        str: 可打印的 TOML
    """
    source_path, dest_path, sources, dest = load_sample_views()
    decisions, _elapsed = align_sample(sources, dest)
    return render_toml(source_path, dest_path, sources, dest, decisions)



def main() -> None:
    """
    函数名: main
    作用: 把样本映射的 TOML 写到标准输出，并把各段耗时写到标准错误
    输入: 无
    输出: 无
    """
    started = time.perf_counter()
    mark = time.perf_counter()
    source_path, dest_path, sources, dest = load_sample_views()
    load_ms = (time.perf_counter() - mark) * 1000
    decisions, align_ms = align_sample(sources, dest)
    mark = time.perf_counter()
    text = render_toml(source_path, dest_path, sources, dest, decisions)
    render_ms = (time.perf_counter() - mark) * 1000
    total_ms = (time.perf_counter() - started) * 1000
    print(text, end="", flush=True)
    print(
        "time_ms"
        f" load={load_ms:.1f}"
        f" map={align_ms['map']:.1f}"
        f" text={align_ms['text']:.1f}"
        f" render={render_ms:.1f}"
        f" total={total_ms:.1f}",
        file=sys.stderr,
        flush=True,
    )



if __name__ == "__main__":
    main()
