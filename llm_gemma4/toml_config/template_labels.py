"""模板标签列表提取。

Phase A: 与 wizard/template_labels.py 完全一致（无 wizard 导入）。
"""

from __future__ import annotations

from pathlib import Path


def list_template_labels(template_path: Path) -> list[str]:
    """
    函数名: list_template_labels
    作用: 从模板 xlsx 首行提取标签列表（与 verify_toml 自洽）
    输入:
        template_path (Path): xlsx 路径
    输出:
        list[str]: 标签名称列表（空表名为 []）
    """
    try:
        from openpyxl import load_workbook
        wb = load_workbook(template_path, read_only=True, data_only=True)
        ws = wb.active
        if ws is None:
            return []
        rows = list(ws.iter_rows(max_row=1, values_only=True))
        labels: list[str] = []
        for row in rows:
            for cell in row:
                if isinstance(cell, str) and cell.strip():
                    labels.append(cell.strip())
        wb.close()
        return labels
    except Exception:
        return []
