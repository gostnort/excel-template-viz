"""根据 WorkflowState 生成合规 TOML：以「当前模板」为底，向导只覆盖已学到的键。

Phase A: 从 wizard/toml_patcher.py 迁移，WorkflowState 来自 llm_gemma4.workflow.state。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.core_toml import (
    TomlGenerator,
    _cell_in_area,
    _core_toml_path,
    _input_area_from_raw,
    _move_to_from_raw,
    _parse_area,
    _parse_input_areas,
    _scan_worksheet_labels_diagonal,
    load_toml,
    offset_cell,
)
from llm_gemma4.workflow.state import FieldState, WorkflowState


def _layout_area_is_set(area: str | list[str] | None) -> bool:
    """
    函数名: _layout_area_is_set
    作用: 判断向导是否已确认 input_area（标量或列表）
    输入:
        area (str | list[str] | None): WorkflowState.input_area
    输出:
        bool: 有非空区域则为 True
    """
    if area is None:
        return False
    if isinstance(area, list):
        return any(str(x).strip() for x in area)
    return bool(str(area).strip())


def _areas_as_list(raw: Any) -> list[str]:
    """
    函数名: _areas_as_list
    作用: 把 input_area 规范成非空字符串列表（供步骤 2 表单预填）
    输入:
        raw (Any): 标量或列表
    输出:
        list[str]: 区域字符串列表；无效则为空列表
    """
    parsed = _input_area_from_raw(raw)
    if parsed is None:
        return []
    if isinstance(parsed, list):
        return [str(x).strip() for x in parsed if str(x).strip()]
    text = str(parsed).strip()
    return [text] if text else []


def _moves_as_list(raw: Any) -> list[str]:
    """
    函数名: _moves_as_list
    作用: 把 move_to 规范成 1～2 个方向列表（供步骤 2 多选预填）
    输入:
        raw (Any): 标量或列表
    输出:
        list[str]: 方向列表；无效则默认 ["down"]
    """
    parsed = _move_to_from_raw(raw)
    if parsed is None:
        return ["down"]
    if isinstance(parsed, list):
        return list(parsed)
    return [str(parsed)]


def _xlsx_sheet_names(template_path: Path) -> tuple[list[str], str]:
    """
    函数名: _xlsx_sheet_names
    作用: 读取模板 xlsx 的工作表名列表与 active 表名
    输入:
        template_path (Path): xlsx 路径
    输出:
        tuple[list[str], str]: (全部表名, active 表名；失败时 ([], ""))
    """
    path = Path(template_path)
    if not path.is_file():
        return [], ""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        names = list(wb.sheetnames)
        active = wb.active.title if wb.active is not None else (names[0] if names else "")
        return names, active
    finally:
        wb.close()


def _ensure_work_sheet(base: dict[str, Any], state: WorkflowState) -> dict[str, Any]:
    """
    函数名: _ensure_work_sheet
    作用: 若底稿 work_sheet 不在本模板 xlsx 中，则改为 CreateDefaultFromTemplate 解析到的真实表名
    输入:
        base (dict[str, Any]): 待写入的配置底稿
        state (WorkflowState): 含 template_path
    输出:
        dict[str, Any]: 校正后的 base（原地修改并返回）
    """
    tpath = state.template_path
    if tpath is None or not Path(tpath).is_file():
        return base
    names, active = _xlsx_sheet_names(Path(tpath))
    if not names:
        return base
    current = str(base.get("work_sheet") or "").strip()
    if current and current in names:
        return base
    # 底稿表名无效：用本模板标准范式重新解析（含首行有标签的表）
    derived = TomlGenerator().CreateDefaultFromTemplate(Path(tpath))
    resolved = str(derived.get("work_sheet") or active or "").strip()
    if not resolved or resolved not in names:
        resolved = active or names[0]
    base["work_sheet"] = resolved
    # 错误表名下的 input_section 不可信；向导未确认 layout 时用本表推导值
    if not _layout_area_is_set(state.input_area):
        section = derived.get("input_section")
        if isinstance(section, dict) and _input_area_from_raw(section.get("input_area")) is not None:
            base["input_section"] = dict(section)
        # fields 若为空则用推导；已有标签列表则保留向导/底稿字段
        if not list(base.get("fields") or []):
            base["fields"] = list(derived.get("fields") or [])
    return base


def _base_config_for_template(state: WorkflowState) -> dict[str, Any]:
    """
    函数名: _base_config_for_template
    作用: 取当前模板的配置底稿——已有 sidecar TOML，否则 CreateDefaultFromTemplate(该 xlsx)
    输入:
        state (WorkflowState): 须含 template_id / template_path
    输出:
        dict[str, Any]: 可序列化配置字典（本模板专属，禁止全局硬编码回退）
    """
    tid = (state.template_id or "").strip()
    base: dict[str, Any] | None = None
    if tid:
        existing = load_toml(tid)
        if existing is not None:
            base = existing.ToDict()
    if base is None:
        tpath = state.template_path
        if tpath is None or not Path(tpath).is_file():
            raise ValueError(
                "cannot build TOML base: need existing templates/{id}/{id}.toml "
                "or a real template_path for CreateDefaultFromTemplate"
            )
        base = TomlGenerator().CreateDefaultFromTemplate(Path(tpath))
    return _ensure_work_sheet(base, state)


def _overlay_wizard_fields(
    base_fields: list[dict[str, Any]],
    wizard_fields: dict[str, FieldState],
) -> list[dict[str, Any]]:
    """
    函数名: _overlay_wizard_fields
    作用: 用向导 FieldState 覆盖底稿 fields（按 Input_label）；底稿没有的标签则追加
    输入:
        base_fields (list[dict]): 模板底稿 [[fields]]
        wizard_fields (dict[str, FieldState]): 向导内存字段
    输出:
        list[dict]: 合并后的 fields 字典列表
    """
    by_label: dict[str, dict[str, Any]] = {}
    order: list[str] = []
    for item in base_fields:
        if not isinstance(item, dict):
            continue
        label = str(item.get("Input_label") or "").strip()
        if not label:
            continue
        by_label[label] = dict(item)
        order.append(label)
    for label, fs in wizard_fields.items():
        row = dict(by_label.get(label) or {
            "Input_label": label,
            "value_from_label": "down",
            "value_offset": 1,
            "field": "",
            "source_file": "",
            "source_sheet": "",
            "index": -1,
            "regex": "",
            "id": False,
        })
        if label not in by_label:
            order.append(label)
        # 向导结果全覆盖底稿（含 index=-1：无文本框列索引）
        row["Input_label"] = label
        if fs.column_name:
            row["field"] = fs.column_name
        if fs.source_file:
            row["source_file"] = fs.source_file
        if fs.source_sheet:
            row["source_sheet"] = fs.source_sheet
        row["index"] = fs.index
        if fs.needs_regex and fs.regex:
            row["regex"] = fs.regex
        elif fs.match_type in ("exact", "fuzzy", "none") or fs.index >= 0:
            # 已跑过匹配：无 regex 时显式空串
            if not fs.needs_regex:
                row["regex"] = ""
        row["id"] = bool(fs.id)
        by_label[label] = row
    return [by_label[label] for label in order if label in by_label]


def _rebuild_fields_for_input_area(
    base: dict[str, Any],
    state: WorkflowState,
    area: str | list[str],
) -> tuple[list[dict[str, Any]], list[str]]:
    """
    函数名: _rebuild_fields_for_input_area
    作用: 按用户确认的 input_area 重建 [[fields]]，只保留值格落在并集内的标签（与 verify 自洽）
    输入:
        base (dict): 当前配置底稿
        state (WorkflowState): 含 template_path
        area (str | list[str]): 用户输入的 input_area（字符串或并集列表）
    输出:
        tuple[list[dict], list[str]]: (保留的 fields, 因越界丢弃的 Input_label)
    """
    tpath = state.template_path
    if tpath is None or not Path(tpath).is_file():
        return list(base.get("fields") or []), []
    try:
        area_rects = _parse_input_areas(area)
    except ValueError:
        try:
            area_rects = [_parse_area(area)]
        except ValueError:
            return list(base.get("fields") or []), []
    work_name = str(base.get("work_sheet") or "").strip() or None
    derived = TomlGenerator().CreateDefaultFromTemplate(Path(tpath), worksheet_name=work_name)
    # 候选来自本模板 xlsx 表头推导；旧 sidecar 仅用于保留仍在区域内字段的已有属性
    base_by_label: dict[str, dict[str, Any]] = {}
    for item in list(base.get("fields") or []):
        if not isinstance(item, dict):
            continue
        label = str(item.get("Input_label") or "").strip()
        if label:
            base_by_label[label] = dict(item)
    candidates: list[dict[str, Any]] = []
    seen: set[str] = set()
    for item in list(derived.get("fields") or []):
        if not isinstance(item, dict):
            continue
        label = str(item.get("Input_label") or "").strip()
        if not label or label in seen:
            continue
        seen.add(label)
        candidates.append(dict(base_by_label.get(label) or item))
    # 底稿里多出的标签（非首行）若值格仍在区域内也保留
    for label, item in base_by_label.items():
        if label in seen:
            continue
        seen.add(label)
        candidates.append(dict(item))
    resolved_sheet = str(
        base.get("work_sheet") or derived.get("work_sheet") or ""
    ).strip()
    wb = load_workbook(Path(tpath), read_only=True, data_only=True)
    try:
        if not resolved_sheet or resolved_sheet not in wb.sheetnames:
            return list(base.get("fields") or []), []
        ws = wb[resolved_sheet]
        label_occurrences = _scan_worksheet_labels_diagonal(ws)
        kept: list[dict[str, Any]] = []
        dropped: list[str] = []
        for item in candidates:
            label = str(item.get("Input_label") or "").strip()
            if not label:
                continue
            occurrences = label_occurrences.get(label) or []
            if not occurrences:
                dropped.append(label)
                continue
            direction = str(item.get("value_from_label") or "down").strip().lower() or "down"
            try:
                voff = int(item.get("value_offset") or 1)
            except (TypeError, ValueError):
                voff = 1
            if voff < 1:
                voff = 1
            # 与 verify_toml 一致：只认值格落入 input_area 并集的候选
            in_area = False
            for label_row, label_col in occurrences:
                try:
                    value_row, value_col = offset_cell(label_row, label_col, direction, voff)
                except ValueError:
                    continue
                if _cell_in_area(value_row, value_col, area_rects):
                    in_area = True
                    break
            if not in_area:
                dropped.append(label)
                continue
            row = dict(item)
            row["Input_label"] = label
            kept.append(row)
        return kept, dropped
    finally:
        wb.close()


def _sync_state_labels_after_layout(
    state: WorkflowState,
    kept_fields: list[dict[str, Any]],
    dropped: list[str],
) -> None:
    """
    函数名: _sync_state_labels_after_layout
    作用: 布局重建后同步 WizardState 的 template_labels / fields（去掉越界标签）
    输入:
        state (WorkflowState): 向导状态
        kept_fields (list[dict]): 重建后保留的 fields
        dropped (list[str]): 丢弃的 Input_label
    输出: 无
    """
    labels = [
        str(item.get("Input_label") or "").strip()
        for item in kept_fields
        if str(item.get("Input_label") or "").strip()
    ]
    state.template_labels = labels
    drop_set = set(dropped)
    for label in list(state.fields.keys()):
        if label in drop_set or (labels and label not in set(labels)):
            del state.fields[label]


def generate_toml(state: WorkflowState) -> str:
    """
    函数名: generate_toml
    作用: 以当前模板底稿合并向导已学状态，生成 TOML 1.0 文本（无全局硬编码骨架）
    输入:
        state (WorkflowState): 向导状态（须能解析到本模板底稿）
    输出:
        str: TOML 文本
    """
    base = _base_config_for_template(state)
    # determiner：预处理完成后写向导结果（含 brace_json 的 ""）；否则保留底稿
    if state.preprocess_done:
        base["determiner"] = "" if state.determiner is None else state.determiner
    # sources：向导已配置则覆盖
    if state.data_sources:
        sources: list[dict[str, Any]] = []
        for source in state.data_sources:
            row = {k: v for k, v in source.items() if k != "type"}
            if row:
                sources.append(row)
        if sources:
            base["sources"] = sources
    # [[input_section]]：用户确认区域后覆盖，并按该区域并集重建 fields（与 verify 自洽）
    if _layout_area_is_set(state.input_area):
        area = state.input_area
        if isinstance(area, list):
            area = [str(x).strip() for x in area if str(x).strip()]
            if len(area) == 1:
                area = area[0]
        else:
            area = str(area).strip()
        move = _move_to_from_raw(state.move_to)
        if move is None:
            # 方向未选时沿用底稿方向
            section = base.get("input_section") or {}
            if isinstance(section, dict):
                move = _move_to_from_raw(section.get("move_to", "down")) or "down"
            else:
                move = "down"
        offset = state.offset if isinstance(state.offset, int) and state.offset >= 1 else None
        if offset is None:
            section = base.get("input_section") or {}
            if isinstance(section, dict):
                try:
                    offset = int(section.get("offset") or 1)
                except (TypeError, ValueError):
                    offset = 1
            else:
                offset = 1
        base["input_section"] = {
            "input_area": area,
            "move_to": move,
            "offset": offset,
        }
        # 禁止沿用旧 sidecar 里越界的 [[fields]]；按 xlsx + 新 area 并集重建
        kept, dropped = _rebuild_fields_for_input_area(base, state, area)
        base["fields"] = kept
        _sync_state_labels_after_layout(state, kept, dropped)
        if state.fields:
            kept_set = {
                str(item.get("Input_label") or "").strip()
                for item in kept
                if str(item.get("Input_label") or "").strip()
            }
            filtered = {
                label: fs for label, fs in state.fields.items() if label in kept_set
            }
            base["fields"] = _overlay_wizard_fields(kept, filtered)
    elif state.fields:
        base_fields = list(base.get("fields") or [])
        base["fields"] = _overlay_wizard_fields(base_fields, state.fields)
    if state.user_inputs.get("db_id_confirmed"):
        db_val = str(state.db_id or "").strip()
        if db_val and db_val not in ("", "None"):
            base["db_id"] = db_val
        else:
            base.pop("db_id", None)
    elif state.db_id:
        base["db_id"] = state.db_id
    return TomlGenerator().ConfigToToml(base)


def persist_wizard_toml(state: WorkflowState, template_id: str = "") -> Path:
    """
    函数名: persist_wizard_toml
    作用: 将合并后的向导 TOML 写入 templates/{id}/{id}.toml
    输入:
        state (WorkflowState): 向导状态
        template_id (str): 模板 ID；空则回退 state.template_id
    输出:
        Path: 写入的 TOML 路径
    """
    tid = (template_id or state.template_id or "").strip()
    if not tid:
        raise ValueError("missing template_id for TOML persist")
    if not state.template_id:
        state.template_id = tid
    path = _core_toml_path(tid)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(generate_toml(state), encoding="utf-8")
    state.written_toml_path = str(path)
    return path


def layout_hints_for_template(state: WorkflowState) -> dict[str, Any]:
    """
    函数名: layout_hints_for_template
    作用: 为步骤 2 表单提供本模板 input_section 提示（保留列表形态；sidecar 优先，否则 xlsx 推导）
    输入:
        state (WorkflowState): 当前向导状态
    输出:
        dict[str, Any]: input_area / move_to / offset（可为 str 或 list）
    """
    def _from_section(section: dict[str, Any]) -> dict[str, Any] | None:
        if not isinstance(section, dict):
            return None
        area = _input_area_from_raw(section.get("input_area"))
        if area is None:
            return None
        move = _move_to_from_raw(section.get("move_to", "down")) or "down"
        try:
            offset = int(section.get("offset") or 1)
        except (TypeError, ValueError):
            offset = 1
        if offset < 1:
            offset = 1
        return {"input_area": area, "move_to": move, "offset": offset}
    # 已有 sidecar：完整 round-trip 列表形态（用户仍可改）
    tid = (state.template_id or "").strip()
    if tid:
        existing = load_toml(tid)
        if existing is not None:
            found = _from_section(dict((existing.ToDict() or {}).get("input_section") or {}))
            if found is not None:
                return found
    # 否则从本模板 xlsx 标准范式推导
    tpath = state.template_path
    if tpath is not None and Path(tpath).is_file():
        try:
            derived = TomlGenerator().CreateDefaultFromTemplate(Path(tpath))
            found = _from_section(dict(derived.get("input_section") or {}))
            if found is not None:
                return found
        except Exception:
            pass
    return {"input_area": "", "move_to": "down", "offset": 1}
