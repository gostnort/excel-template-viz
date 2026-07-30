"""从模板 xlsx 扫描 Input_label 列表（不依赖已有 TOML）。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.core_toml import _scan_worksheet_labels_diagonal


def list_template_labels(template_path: Path, *, work_sheet: str | None = None) -> list[str]:
    """
    函数名: list_template_labels
    作用: 对角扫描工作表标签文本，返回去重后的 label 列表（按扫描顺序）
    输入:
        template_path (Path): 模板 xlsx 路径
        work_sheet (str | None): 工作表名，None 时用 active sheet
    输出:
        list[str]: Input_label 候选列表
    """
    wb = load_workbook(template_path, read_only=True, data_only=True)
    try:
        if work_sheet and work_sheet in wb.sheetnames:
            ws = wb[work_sheet]
        else:
            ws = wb.active
        label_map, _dupes = _scan_worksheet_labels_diagonal(ws)
        return list(label_map.keys())
    finally:
        wb.close()
