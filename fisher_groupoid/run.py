"""调用脚本：指定文件、sheet、哪些格子是 Label、纤维方向，再交给引擎。"""

from __future__ import annotations

import argparse
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Sequence

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, range_boundaries

from fisher_groupoid.engine import (
    ChartAtlas,
    FiberSeries,
    GeodesicEngine,
    OutputGroup,
    excel_coord,
    infer_morphism_phi,
    infer_object_key,
    output_group_from_header,
    records_from_fibers,
    slot_cell,
)
from fisher_groupoid.meaning import (
    ensure_labels_resolved,
    key_map_to_vector,
    set_clarifier,
    set_meaning_open_fn,
    set_probe_scorer,
    space_dim,
)
from llm_lmstudio.chat import CHAT_TIMEOUT, _extract_text
from llm_lmstudio.client import LmStudioError, request_json
from llm_lmstudio.json_extract import extract_json_object
from llm_lmstudio.models import LOAD_TIMEOUT, list_models, remember_loaded, reset_model_session


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = PROJECT_ROOT / "docs" / "sample" / "执法堂业绩.xlsx"
DEFAULT_TEMPLATE = PROJECT_ROOT / "docs" / "sample" / "sample_template.xlsx"
DEFAULT_OUT = PROJECT_ROOT / "plans" / "geodesic_stage_out.json"
TEST_LM_MODEL = "gemma-4-e4b-it"
TEST_LM_LOAD = {"flash_attention": True, "context_length": 2048}
_lm_model: str = ""
_lm_load: dict = {}
_lm_ready: bool = False


@dataclass(frozen=True)
class ChartRequest:
    """一张源图：声明哪些格子是 Label，纤维沿 direction 走。"""

    sheet: str
    label_cells: tuple[str, ...]
    direction: str = "down"




@dataclass(frozen=True)
class TemplateDecl:
    """输出表：指定 Label 行，旁侧填写为纤维。"""

    workbook: Path
    sheet: str = "Input_sheet"
    label_row: int = 1
    phi: tuple[tuple[str, str], ...] = ()



@dataclass(frozen=True)
class CollectJob:
    """一次采集任务：文件 + 多张图 + 对象键。"""

    workbook: Path
    charts: tuple[ChartRequest, ...]
    out_path: Path
    object_key: str = ""
    lm_model: str = TEST_LM_MODEL
    template: TemplateDecl | None = None


def default_job() -> CollectJob:
    """
    函数名: default_job
    作用: 执法堂样本的采集声明（只在调用层）
    输入: 无
    输出:
        CollectJob
    """
    return CollectJob(
        workbook=DEFAULT_WORKBOOK,
        charts=(
            ChartRequest("sheet1", ("A1:D1",), "down"),
            ChartRequest("sheet2", ("A1:C1",), "down"),
        ),
        object_key="",
        out_path=DEFAULT_OUT,
        template=TemplateDecl(workbook=DEFAULT_TEMPLATE, label_row=1),
    )


def bind_lm(model: str | None = None) -> None:
    """
    函数名: bind_lm
    作用: 绑定 LM Studio 模型，并把原生向量函数交给意义层
    输入:
        model (str | None): 空则用 TEST_LM_MODEL
    输出: 无
    """
    global _lm_model, _lm_load, _lm_ready
    name = str(model or TEST_LM_MODEL).strip()
    if not name:
        raise ValueError("模型名不能为空")
    if name != _lm_model:
        _lm_ready = False
    _lm_model = name
    _lm_load = dict(TEST_LM_LOAD)
    set_probe_scorer(score_model_vector, cache_key=name, persist=True)
    set_meaning_open_fn(header_meaning_open)
    set_clarifier(request_user_clarification)


def _handshake() -> None:
    """
    函数名: _handshake
    作用: 只在首次（或断线后）确认模型已加载；之后只做前向
    输入: 无
    输出: 无
    """
    global _lm_ready
    if _lm_ready:
        return
    key = _lm_model
    loaded = False
    # 中文注释: 只 list 一次，不按 key 打 GET /models/{id}
    for item in list_models():
        if str(item.get("key") or item.get("id") or "") != key:
            continue
        instances = item.get("loaded_instances") or []
        loaded = isinstance(instances, list) and len(instances) > 0
        break
    if not loaded:
        body = {"model": key}
        body.update(_lm_load)
        request_json(
            "POST",
            "/api/v1/models/load",
            body=body,
            timeout=LOAD_TIMEOUT,
        )
    remember_loaded(key)
    _lm_ready = True


def _parse_embedding_payload(raw: Any) -> np.ndarray | None:
    """
    函数名: _parse_embedding_payload
    作用: 从 embeddings 响应取出完整向量，不裁切
    输入:
        raw (Any): HTTP JSON
    输出:
        np.ndarray | None
    """
    if not isinstance(raw, dict):
        return None
    data = raw.get("data")
    if isinstance(data, list) and data:
        first = data[0]
        if isinstance(first, dict):
            emb = first.get("embedding")
            if isinstance(emb, list) and emb:
                return np.array([float(x) for x in emb], dtype=np.float64)
            if isinstance(emb, dict) and emb:
                try:
                    return key_map_to_vector(emb)
                except ValueError:
                    return None
    emb = raw.get("embedding")
    if isinstance(emb, list) and emb:
        return np.array([float(x) for x in emb], dtype=np.float64)
    return None


def score_model_vector(text: str) -> np.ndarray:
    """
    函数名: score_model_vector
    作用: 用已加载模型对整句做 token 前向，返回原生维数向量
    输入:
        text (str)
    输出:
        np.ndarray
    """
    global _lm_ready
    _handshake()
    last_error = "no embedding"
    attempts = (
        ("/v1/embeddings", {"model": _lm_model, "input": text}),
        ("/api/v0/embeddings", {"model": _lm_model, "input": text}),
        ("/api/v1/embeddings", {"model": _lm_model, "input": text}),
    )
    for path, body in attempts:
        try:
            raw = request_json("POST", path, body=body, timeout=CHAT_TIMEOUT)
        except LmStudioError as exc:
            last_error = str(exc)
            _lm_ready = False
            reset_model_session()
            if exc.status is None or (exc.status is not None and exc.status >= 500):
                _handshake()
            continue
        vec = _parse_embedding_payload(raw)
        if vec is None or vec.size < 1:
            last_error = "empty embedding"
            continue
        return vec
    raise RuntimeError(f"LM 未返回模型维数向量: {last_error}")


def header_meaning_open(label: str, vector: np.ndarray) -> bool:
    """
    函数名: header_meaning_open
    作用: 问已加载模型该表头含义是否唯一；不唯一或无法判定则必须问人
    输入:
        label (str)
        vector (np.ndarray): 已采样向量，此处不裁切
    输出:
        bool: True 表示必须问人
    """
    _ = vector
    _handshake()
    user = f"Spreadsheet column header:\n{label}\n"
    last_error = "no JSON"
    for _ in range(2):
        try:
            raw = request_json(
                "POST",
                "/api/v1/chat",
                body={
                    "model": _lm_model,
                    "input": user,
                    "temperature": 0.0,
                    "store": False,
                    "system_prompt": (
                        "Judge this spreadsheet header as a data-field chart. "
                        'Default {"unique": true}. '
                        "Set unique false only if the header is underspecified as a field: "
                        "an acronym, abbreviation, ALL-CAPS token, or a form that is not "
                        "a complete ordinary word, so a reader would have to guess among "
                        "incompatible fields. "
                        "Do not set unique false merely because an ordinary complete word "
                        "has extra senses in a dictionary. "
                        'Reply with one JSON object only: {"unique": true} or {"unique": false}. '
                        "No markdown."
                    ),
                    "max_output_tokens": 64,
                },
                timeout=CHAT_TIMEOUT,
            )
        except LmStudioError as exc:
            last_error = str(exc)
            break
        text_out = ""
        if isinstance(raw, dict):
            text_out, _, _ = _extract_text(raw)
        payload, err = extract_json_object(text_out)
        if err or payload is None:
            last_error = err or "empty"
            continue
        if "unique" not in payload:
            last_error = "missing unique"
            continue
        flag = payload.get("unique")
        if flag is True:
            return False
        if flag is False:
            return True
        last_error = "unique not bool"
    raise RuntimeError(f"模型未能判定表头含义是否唯一: {last_error}")


def request_user_clarification(label_raw: str, candidates: list[str]) -> str:
    """
    函数名: request_user_clarification
    作用: 控制台询问 Label 释义
    输入:
        label_raw (str): 原始 Label
        candidates (list[str]): 坐标键候选
    输出:
        str: 用户给出的完全体释义
    """
    prompt = f"这个列名是什么意思？ Label={label_raw!r} 请给出完整含义: "
    text = input(prompt).strip()
    if not text:
        raise ValueError("澄清释义不能为空")
    return text


def request_section_break(
    unfilled: Sequence[str],
    unpushed: Sequence[str],
) -> dict[str, str]:
    """
    函数名: request_section_break
    作用: 输出底点不在输入群胚内时向人确认破缺，不得自行补列
    输入:
        unfilled (Sequence[str]): 未覆盖的输出 Label
        unpushed (Sequence[str]): 未推出的源 Label
    输出:
        dict[str, str]: 人指定的 输出 Label -> 源 Label；未指定的保持破缺
    """
    dests = [str(name).strip() for name in unfilled if str(name).strip()]
    sources = [str(name).strip() for name in unpushed if str(name).strip()]
    print("群破缺：下列输出底点不在输入群胚内，不会强行补足。")
    print("未覆盖输出: " + (", ".join(dests) if dests else "（无）"))
    print("未推出源: " + (", ".join(sources) if sources else "（无）"))
    assigned: dict[str, str] = {}
    allowed = {name: name for name in sources}
    for dest in dests:
        prompt = (
            f"输出列 {dest!r} 在输入群胚中没有对应。"
            "源里没有请留空；若应对某个未推出源列，请输入源列名: "
        )
        while True:
            text = input(prompt).strip()
            if not text:
                break
            key = text.casefold()
            match = next((name for name in allowed if name.casefold() == key), "")
            if match:
                assigned[dest] = match
                del allowed[match]
                break
            print(f"源列 {text!r} 不在未推出集合中，请重输或留空。")
    return assigned


def parse_cell(addr: str) -> tuple[int, int]:
    """
    函数名: parse_cell
    作用: Excel 地址 -> (行, 列)，均为 1-based
    输入:
        addr (str): 如 A1
    输出:
        tuple[int, int]
    """
    col_letter, row = coordinate_from_string(addr.strip().upper())
    return int(row), int(column_index_from_string(col_letter))


def expand_cell_specs(specs: tuple[str, ...]) -> list[str]:
    """
    函数名: expand_cell_specs
    作用: 把 A1 或 A1:E1 展开成单元格列表，保留声明顺序
    输入:
        specs (tuple[str, ...])
    输出:
        list[str]
    """
    cells: list[str] = []
    seen: set[str] = set()
    for spec in specs:
        token = spec.strip().upper()
        min_col, min_row, max_col, max_row = range_boundaries(token)
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                addr = f"{get_column_letter(col)}{row}"
                if addr in seen:
                    continue
                seen.add(addr)
                cells.append(addr)
    return cells


def parse_chart_spec(spec: str) -> ChartRequest:
    """
    函数名: parse_chart_spec
    作用: 解析 sheet:A1:D1 或 sheet:A1,B1 为 Label 声明
    输入:
        spec (str)
    输出:
        ChartRequest
    """
    sheet, sep, rest = spec.partition(":")
    if not sep or not sheet.strip() or not rest.strip():
        raise ValueError(f"chart 格式应为 sheet:A1:D1 或 sheet:A1,B1，收到 {spec!r}")
    token = rest.strip()
    direction = "down"
    if "|" in token:
        cells_part, mark = token.rsplit("|", 1)
        token = cells_part.strip()
        mark = mark.strip().lower()
        if mark not in {"down", "right"}:
            raise ValueError(f"纤维方向只能是 down 或 right，收到 {mark!r}")
        direction = mark
    cells = tuple(part.strip() for part in token.split(",") if part.strip())
    if not cells:
        raise ValueError(f"chart 未给出 Label 格子: {spec!r}")
    return ChartRequest(sheet=sheet.strip(), label_cells=cells, direction=direction)


def cell_text(value: Any) -> str:
    """
    函数名: cell_text
    作用: 单元格转字符串；空 / Null 为未感知
    输入:
        value (Any)
    输出:
        str
    """
    if value is None:
        return ""
    if isinstance(value, float) and value != value:
        return ""
    text = str(value).strip()
    if text.lower() in {"none", "null", "nan"}:
        return ""
    return text


def collect_chart(path: Path, request: ChartRequest) -> tuple[list[FiberSeries], str]:
    """
    函数名: collect_chart
    作用: 按声明的 Label 格子读纤维；空格、隔行、不相邻列都只是缺位置
    输入:
        path (Path): xlsx
        request (ChartRequest)
    输出:
        tuple[list[FiberSeries], str]: 纤维与建议对象键
    """
    if request.direction not in {"down", "right"}:
        raise ValueError(f"纤维方向只能是 down 或 right，收到 {request.direction!r}")
    workbook = load_workbook(path, data_only=True)
    try:
        if request.sheet not in workbook.sheetnames:
            raise KeyError(request.sheet)
        worksheet = workbook[request.sheet]
        addrs = expand_cell_specs(request.label_cells)
        label_at: dict[tuple[int, int], str] = {}
        order: list[tuple[int, int]] = []
        # 中文注释: 空单元格不是 Label；不要求同一行/同一列
        for addr in addrs:
            row_idx, col_idx = parse_cell(addr)
            name = cell_text(worksheet.cell(row_idx, col_idx).value)
            if not name:
                continue
            key = (row_idx, col_idx)
            if key not in label_at:
                order.append(key)
            label_at[key] = name
        if not label_at:
            raise ValueError(f"{request.sheet} 声明的格子里没有读到 Label")
        names = [label_at[key] for key in order]
        ensure_labels_resolved(names)
        occupied = set(label_at)
        max_row = int(worksheet.max_row or 1)
        max_col = int(worksheet.max_column or 1)
        fibers: list[FiberSeries] = []
        for row_idx, col_idx in order:
            name = label_at[(row_idx, col_idx)]
            at: dict[int, str] = {}
            if request.direction == "down":
                for data_row in range(row_idx + 1, max_row + 1):
                    if (data_row, col_idx) in occupied:
                        continue
                    at[data_row] = cell_text(worksheet.cell(data_row, col_idx).value)
            else:
                for data_col in range(col_idx + 1, max_col + 1):
                    if (row_idx, data_col) in occupied:
                        continue
                    at[data_col] = cell_text(worksheet.cell(row_idx, data_col).value)
            fibers.append(FiberSeries(label=name, source=request.sheet, at=at))
        return fibers, infer_object_key(names)
    finally:
        workbook.close()


def collect_atlases(job: CollectJob) -> list[ChartAtlas]:
    """
    函数名: collect_atlases
    作用: 每张源表自成一图
    输入:
        job (CollectJob)
    输出:
        list[ChartAtlas]
    """
    if not job.workbook.is_file():
        raise FileNotFoundError(job.workbook)
    atlases: list[ChartAtlas] = []
    for chart in job.charts:
        fibers, suggested = collect_chart(job.workbook, chart)
        object_key = str(job.object_key or suggested)
        labels, records = records_from_fibers(fibers, object_key)
        atlases.append(
            ChartAtlas(
                name=chart.sheet,
                labels=tuple(labels),
                records=tuple(records),
            )
        )
    return atlases


def collect_output_samples(decl: TemplateDecl) -> tuple[OutputGroup, list[dict]]:
    """
    函数名: collect_output_samples
    作用: 读指定 Label 行与其下填写，交给引擎建输出群
    输入:
        decl (TemplateDecl)
    输出:
        tuple[OutputGroup, list[dict]]: 群对象与已填样本截面
    """
    if not decl.workbook.is_file():
        raise FileNotFoundError(decl.workbook)
    workbook = load_workbook(decl.workbook, data_only=True)
    try:
        if decl.sheet not in workbook.sheetnames:
            raise KeyError(decl.sheet)
        worksheet = workbook[decl.sheet]
        label_row = int(decl.label_row)
        max_col = int(worksheet.max_column or 1)
        max_row = int(worksheet.max_row or 1)
        columns: list[tuple[int, str]] = []
        for col in range(1, max_col + 1):
            name = cell_text(worksheet.cell(label_row, col).value)
            if name:
                columns.append((col, name))
        names = [name for _col, name in columns]
        ensure_labels_resolved(names)
        group = output_group_from_header(label_row, columns, fiber_down=1)
        samples: list[dict] = []
        # 中文注释: Label 行之下每个非空行是 G 上一个已填样本，空行不算点
        for data_row in range(label_row + 1, max_row + 1):
            k = data_row - (label_row + 1)
            sigma: dict[str, str] = {}
            cells: list[dict] = []
            empty = True
            for anchor in group.anchors:
                row_idx, col_idx = slot_cell(group, anchor, k, 0, fiber=True)
                text = cell_text(worksheet.cell(row_idx, col_idx).value)
                sigma[anchor.label] = text
                if text:
                    empty = False
                cells.append(
                    {
                        "cell": excel_coord(row_idx, col_idx),
                        "label": anchor.label,
                        "literal": text,
                    }
                )
            if empty:
                continue
            samples.append({"g": {"k": k, "m": 0}, "sigma_template": sigma, "cells": cells})
        return group, samples
    finally:
        workbook.close()


def job_meta(job: CollectJob) -> dict:
    """
    函数名: job_meta
    作用: 采集声明写入报告，与引擎几何分开
    输入:
        job (CollectJob)
    输出:
        dict
    """
    payload = {
        "workbook": str(job.workbook),
        "object_key": job.object_key,
        "lm_model": job.lm_model,
        "charts": [
            {
                "sheet": chart.sheet,
                "label_cells": list(chart.label_cells),
                "direction": chart.direction,
            }
            for chart in job.charts
        ],
    }
    if job.template is not None:
        payload["template"] = {
            "workbook": str(job.template.workbook),
            "sheet": job.template.sheet,
            "label_row": job.template.label_row,
            "phi": [{"source": src, "output": dst} for src, dst in job.template.phi],
        }
    return payload


def build_job(argv: list[str] | None = None) -> CollectJob:
    """
    函数名: build_job
    作用: 命令行覆盖默认采集声明
    输入:
        argv (list[str] | None)
    输出:
        CollectJob
    """
    parser = argparse.ArgumentParser(description="采集 xlsx 后交给 GOP 引擎")
    parser.add_argument("--file", type=Path, default=None, help="xlsx 路径")
    parser.add_argument(
        "--chart",
        action="append",
        default=None,
        help="sheet:A1:D1 或 sheet:A1,B1 声明 Label 格，可重复；可加 |down 或 |right",
    )
    parser.add_argument("--key", default=None, help="对象键，对应某个 Label 名")
    parser.add_argument("--model", default=None, help="LM Studio 模型 key，仅调用层")
    parser.add_argument("--out", type=Path, default=None, help="JSON 输出路径")
    args = parser.parse_args(argv)
    if args.chart:
        if args.file is None:
            parser.error("--chart 需要同时给 --file")
        charts = tuple(parse_chart_spec(item) for item in args.chart)
        return CollectJob(
            workbook=args.file,
            charts=charts,
            object_key=args.key or "",
            out_path=args.out or DEFAULT_OUT,
            lm_model=args.model or TEST_LM_MODEL,
            template=default_job().template,
        )
    job = default_job()
    return CollectJob(
        workbook=args.file or job.workbook,
        charts=job.charts,
        object_key=args.key or job.object_key,
        out_path=args.out or job.out_path,
        lm_model=args.model or TEST_LM_MODEL,
        template=job.template,
    )


def run(job: CollectJob | None = None) -> Path:
    """
    函数名: run
    作用: 采集 -> 引擎 -> 写 JSON
    输入:
        job (CollectJob | None): 空则用默认样本声明
    输出:
        Path: 报告路径
    """
    loaded = job or default_job()
    bind_lm(loaded.lm_model)
    engine = GeodesicEngine()
    atlases = collect_atlases(loaded)
    report = engine.input_groupoid(atlases)
    report["model"] = _lm_model
    report["dim"] = space_dim()
    report["collect"] = job_meta(loaded)
    if loaded.template is not None:
        group, samples = collect_output_samples(loaded.template)
        if loaded.template.phi:
            phi = dict(loaded.template.phi)
        else:
            source_names: list[str] = []
            for obj in report["objects"]:
                for name in obj.get("sigma_in") or {}:
                    if name not in source_names:
                        source_names.append(str(name))
            dest_names = [anchor.label for anchor in group.anchors]
            residue_samples = [
                dict(obj.get("sigma_in") or {})
                for obj in report["objects"]
                if obj.get("morphism_product") == 1
            ]
            phi = infer_morphism_phi(source_names, dest_names, residue_samples)
        ready = [obj for obj in report["objects"] if obj.get("pushforward_ready")]
        isolated = [obj for obj in report["objects"] if not obj.get("pushforward_ready")]
        report["output_group"] = engine.push_to_output_group(
            group,
            samples,
            ready,
            phi,
            isolated=isolated,
        )
        out = report["output_group"]
        if out.get("section_broken"):
            extra = request_section_break(
                out.get("unfilled_output_labels") or [],
                out.get("unpushed_source_labels") or [],
            )
            report["section_break_user"] = extra
            if extra:
                for dst, src in extra.items():
                    phi[src] = dst
                report["output_group"] = engine.push_to_output_group(
                    group,
                    samples,
                    ready,
                    phi,
                    isolated=isolated,
                )
    loaded.out_path.parent.mkdir(parents=True, exist_ok=True)
    loaded.out_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return loaded.out_path


def main() -> None:
    """
    函数名: main
    作用: CLI 入口
    输入: 无
    输出: 无
    """
    path = run(build_job())
    print(str(path))


if __name__ == "__main__":
    main()
